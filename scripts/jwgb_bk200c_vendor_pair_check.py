"""End-to-end check: production JWGBT import vs a same-run vendor Excel export.

Usage (PowerShell, repository root):
  python -m scripts.jwgb_bk200c_vendor_pair_check "<run>.jwgbt" "<export>.xlsx" --output report.json

Reads the vendor workbook's Isotherm sheet (ID, p, p/p0, V, p0, Time; branches
separated by a '---' row, desorption starting with the peak reference), pairs
rows to the imported points by branch, time and pressure, and compares the
production quantity to export rounding (half a unit of the 5th decimal).
A pairing mismatch raises instead of comparing unrelated runs. Passing this for
one file is consistency with that export, not proof of physical accuracy.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path

from tristar_bet.jwgb_native import load_jwgb_native

HALF_UNIT = 0.000005 + 1e-10


def read_export(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [list(r) for r in wb['Isotherm'].iter_rows(values_only=True)]
    finally:
        wb.close()
    header, branches, current = rows[0], [[]], None
    if [str(h).split(' ')[0] for h in header[:3]] != ['ID', 'p', 'p/p0']:
        raise ValueError(f'Unexpected Isotherm header: {header}')
    for row in rows[1:]:
        if row[0] is None:
            continue
        if str(row[0]).strip().startswith('---'):
            branches.append([])
            continue
        branches[-1].append({'id': row[0], 'pressure_kpa': float(row[1]), 'relative_pressure': float(row[2]),
                             'v': float(row[3]), 'p0': float(row[4]), 'time': str(row[5])})
    if len(branches) != 2:
        raise ValueError('Expected adsorption and desorption branches separated by ---')
    return branches


def check(native_path, export_path):
    result = load_jwgb_native(native_path)
    options = result.method_options
    if not options.get('jwgb_reconstructed'):
        raise ValueError(f"Import produced no quantities: {options.get('jwgb_reconstruction_unavailable_reason')}")
    ads_export, des_export = read_export(export_path)
    start = datetime.fromisoformat(options['test_started_time'])
    rows = {'adsorption': ads_export, 'desorption': des_export[1:]}  # drop peak reference
    report_rows, summary = [], {}
    for phase, export in rows.items():
        points = [p for p in result.isotherm if p.phase == phase]
        if len(points) != len(export):
            raise ValueError(f'{phase}: {len(points)} imported vs {len(export)} exported points')
        hits, worst = 0, 0.0
        for point, ref in zip(points, export):
            kpa = point.absolute_pressure_mmHg * 101.325 / 760
            elapsed = (datetime.fromisoformat(f"{start.date()}T{ref['time']}") - start).total_seconds()
            if abs(kpa - ref['pressure_kpa']) > 5.01e-9 or abs(point.relative_pressure - ref['relative_pressure']) > 5.01e-9 \
                    or (elapsed >= 0 and abs(elapsed - point.elapsed_seconds) > 1):
                raise ValueError(f'{phase} point {ref["id"]} does not pair with the export (pressure/time)')
            error = point.quantity_adsorbed_cm3_g_stp - ref['v']
            ok = abs(error) <= HALF_UNIT
            hits += ok
            worst = max(worst, abs(error))
            report_rows.append({'phase': phase, 'export_id': ref['id'], 'relative_pressure': point.relative_pressure,
                                'import_cm3_g': point.quantity_adsorbed_cm3_g_stp, 'export_cm3_g': ref['v'],
                                'error_cm3_g': error, 'matches_export_precision': ok})
        summary[phase] = {'points': len(export), 'matching': hits, 'max_abs_error_cm3_g': worst}
    digest = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
    return {
        'native': str(Path(native_path).resolve()), 'native_sha256': digest(native_path),
        'export': str(Path(export_path).resolve()), 'export_sha256': digest(export_path),
        'profile': options.get('jwgb_reconstruction_profile'),
        'nonideality_applied': options.get('jwgb_reconstruction_nonideality_applied'),
        'closure': options.get('jwgb_reconstruction_closure'),
        'summary': summary,
        'all_points_match': all(s['matching'] == s['points'] for s in summary.values()),
        'rows': report_rows,
        'limitation': 'Consistency with one vendor export; not physical accuracy or validation of other runs.',
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('native')
    parser.add_argument('export')
    parser.add_argument('--output')
    args = parser.parse_args()
    report = check(args.native, args.export)
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(text, encoding='utf-8')
    print(json.dumps({'summary': report['summary'], 'all_points_match': report['all_points_match']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
