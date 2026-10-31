"""Independently audit BET fits in JWGB Excel exports (read-only).

Use the full isotherm precision and the IDs explicitly listed in the BET table.
This validates the uptake-to-BET stage, never the pressure-to-uptake stage.
Run: python -m scripts.jwgb_export_bet_audit FILE [FILE ...] --output REPORT.json
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path

from scripts.jwgb_paired_probe import bet_hypothesis, bet_diagnostics


def _decimal(value):
    # Observed German exports use a decimal comma without thousands separators.
    return float(str(value).strip().replace(',', '.'))


def _table(sheet, required):
    """Locate observed named columns rather than assume a fixed header row."""
    header_row = None
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), sheet.min_row):
        labels = [str(value).strip().replace('³', '3') for value in row]
        if all(label in labels for label in required):
            indices = [labels.index(label) for label in required]
            header_row = row_number
            break
    else:
        raise ValueError(f'Missing columns in {sheet.title}: {required}')
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if len(row) <= max(indices):
            continue
        try:
            values = [_decimal(row[i]) for i in indices]
        except (TypeError, ValueError) as exc:
            # A separator ends the adsorption branch, once data have begun.
            if any(str(v).strip() == '---' for v in row):
                break
            if all(value is None or not str(value).strip() for value in row):
                continue
            raise ValueError(f'Malformed data row in {sheet.title}') from exc
        if not all(math.isfinite(v) for v in values):
            raise ValueError('Non-finite export value')
        if not values[0].is_integer():
            raise ValueError('Non-integer measurement ID')
        yield {'phase': 'adsorption', 'measurement': int(values[0]),
               'relative_pressure': values[1], 'candidate_cm3_g': values[2]}


def read_export(path):
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        required = ('ID', 'p/p0', 'V (cm3/g STP)')
        rows = list(_table(workbook['Isotherm'], required))
        selected = list(_table(workbook['BET Surface Area'], required))
        area = None
        for name in ('Info', 'Summary'):
            if name not in workbook.sheetnames:
                continue
            for row in workbook[name].iter_rows(values_only=True):
                for i, cell in enumerate(row[:-1]):
                    if str(cell).strip().rstrip(':') == 'BET Surface Area':
                        area = _decimal(row[i + 1])
        if area is None or not math.isfinite(area) or area <= 0:
            raise ValueError('Missing positive report BET area')
        return {'source_file': path.name,
                'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'adsorption_rows': rows, 'bet_table_rows': selected,
                'reported_area_m2_g': area}
    finally:
        workbook.close()


def audit_export(data, *, report_decimals=3):
    rows, selected = data['adsorption_rows'], data['bet_table_rows']
    by_id = {row['measurement']: row for row in rows}
    if len(by_id) != len(rows):
        raise ValueError('Duplicate adsorption IDs')
    ids = [row['measurement'] for row in selected]
    for row in selected:
        original = by_id.get(row['measurement'])
        if original is None:
            raise ValueError('BET table ID absent from adsorption branch')
        # Both observed BET table columns use five decimals; isotherm x uses eight.
        if (abs(row['relative_pressure'] - original['relative_pressure']) > 5.01e-6
                or abs(row['candidate_cm3_g'] - original['candidate_cm3_g']) > 5.01e-6):
            raise ValueError('BET table does not match the isotherm')
    fit = bet_hypothesis(rows, ids)
    rounded_fit = bet_hypothesis(selected, ids)
    expected = float(data['reported_area_m2_g'])
    half_unit = 0.5 * 10 ** (-report_decimals)
    modern_area = fit['vm_cm3_g'] * (6.02214076e23 * .162e-18 / 22414)
    return {
        'source_file': data['source_file'], 'source_sha256': data['source_sha256'],
        'validation_scope': 'exported uptake to BET only; no raw-pressure reconstruction validation',
        'raw_reconstruction_validated': False,
        'precision_note': 'Exported input rounding is not an uncertainty bound; exact native values may give a different last digit.',
        'fit': fit, 'reported_area_m2_g': expected,
        'area_error_m2_g': fit['surface_area_m2_g'] - expected,
        'matches_report_rounding': abs(fit['surface_area_m2_g'] - expected) <= half_unit,
        'report_decimals': report_decimals,
        'controls': {
            'fit_using_rounded_bet_table_area': rounded_fit['surface_area_m2_g'],
            'fit_using_rounded_bet_table_error': rounded_fit['surface_area_m2_g'] - expected,
            'modern_constants_area': modern_area,
            'modern_constants_error': modern_area - expected,
        },
        'diagnostics': bet_diagnostics(rows, fit),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('files', nargs='+', type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--report-decimals', type=int, default=3, choices=range(1, 9))
    args = parser.parse_args()
    if args.output.resolve() in {path.resolve() for path in args.files}:
        parser.error('Output must not overwrite a source')
    reports = [audit_export(read_export(path), report_decimals=args.report_decimals) for path in args.files]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(reports, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    for report in reports:
        print(f"{report['source_file']}: {report['fit']['surface_area_m2_g']:.9f}; "
              f"matches report rounding: {report['matches_report_rounding']}")


if __name__ == '__main__':
    main()
