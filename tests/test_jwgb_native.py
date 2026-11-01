import os
import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from tristar_bet import load_file, analysis_bundle, export_results_csv
from tristar_bet.jwgb_native import JwgbNativeParseError
from tristar_bet.analysis import bet_analysis

XML = '''<?xml version="1.0" encoding="UTF-8"?>
<doc Ver="1">
<General Ver="1"><GeneralConfig><AdsorbGas>N2</AdsorbGas><AdsorbTemp>77.35</AdsorbTemp></GeneralConfig></General>
<StationConfig Ver="1"><Station><Model>BK200C</Model></Station></StationConfig>
<Config Ver="1"><Base><PressUnits>0</PressUnits><Version>00.01.05.00</Version></Base>
<SampleInfo><SampleName></SampleName><SampleNum>1</SampleNum><SampleWeight>0.10784</SampleWeight></SampleInfo>
<Vd>30</Vd><QValue>0.0085</QValue></Config>
<Data Ver="1"><BaseInfo><StartTime>2026-10-12 11:00:00</StartTime><EndTime>2026-10-12 12:00:00</EndTime></BaseInfo>
<Adsorb>
<Item ID="0" Pd="-0.00001" Pcd="0.001" P0="100" Day="2026-10-12" Time="11:01:00"/>
<Item ID="1" Pd="20" Pcd="10" P0="100" Day="2026-10-12" Time="11:10:00"/>
<Item ID="2" Pd="30" Pcd="20" P0="100" Day="2026-10-12" Time="11:20:00"/>
<Item ID="3" Pd="40" Pcd="30" P0="100" Day="2026-10-12" Time="11:30:00"/>
</Adsorb><Doff>
<Item ID="1" Pd="-0.00001" Pcd="30" P0="100" Day="2026-10-12" Time="11:31:00"/>
<Item ID="2" Pd="10" Pcd="20" P0="100" Day="2026-10-12" Time="11:40:00"/>
</Doff></Data><Log Ver="1"><Log><Item Date="2026.10.12" Log="Test complete!"/></Log></Log></doc>'''


def reconstructable_xml():
    return (XML.replace('</GeneralConfig>', '<CrossSectionalArea>0.162</CrossSectionalArea>'
                        '<IsRealGas>0</IsRealGas><IsInsitu>0</IsInsitu>'
                        '<NonIdealityFactor>0.00049467</NonIdealityFactor></GeneralConfig>')
            .replace('</Base>', '<DoseGasType>0</DoseGasType></Base>')
            .replace('</Config>', '<AutoTestQ>0</AutoTestQ><FastTestQ>0</FastTestQ>'
                     '<IsMicroproe>0</IsMicroproe><MultiBETReport><MinRange>0.08</MinRange>'
                     '<MaxRange>0.30</MaxRange></MultiBETReport></Config>')
            .replace(' Day=', ' PdT="26.85" PcdT="26.85" Balance="1" SubIndex="0" Day='))


class NativeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "sample.JWGBT"
        self.path.write_text(XML, encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def test_pressure_conversion_and_missing_uptake(self):
        result = load_file(self.path)
        self.assertEqual(result.sample_name, "1")
        self.assertEqual(result.point_count, 4)
        self.assertEqual(len(result.method_options["jwgb_native_pressure_records"]), 6)
        self.assertEqual(result.test_duration_seconds, 3600)
        self.assertAlmostEqual(result.isotherm[0].relative_pressure, 0.1)
        self.assertAlmostEqual(result.isotherm[0].absolute_pressure_mmHg, 10 * 760 / 101.325)
        self.assertEqual(result.isotherm[-1].phase, "desorption")
        self.assertEqual(result.isotherm[0].elapsed_seconds, 600)
        for point in result.isotherm:
            self.assertIsNone(point.raw_internal_cm3_stp)
            self.assertIsNone(point.quantity_adsorbed_cm3_g_stp)
        for fit in analysis_bundle(result).values():
            self.assertFalse(fit.ok)
            self.assertIsNone(fit.surface_area_m2_g)
        self.assertEqual(bet_analysis(result).status, 'native_uptake_unavailable')
        export_results_csv([result], Path(self.temp.name) / "export")

    def test_rejects_bad_documents_and_units(self):
        for value in (XML[:-10], XML.replace('Ver="1"', 'Ver="2"', 1),
                      XML.replace('<PressUnits>0', '<PressUnits>1'),
                      XML.replace('P0="100"', 'P0="0"', 1),
                      XML.replace('Pcd="10"', 'Pcd="NaN"'),
                      XML.replace('ID="3"', 'ID="2"'),
                      XML.replace('ID="1" Pd="20"', 'ID="1_1" SubIndex="2" Pd="20"'),
                      XML.replace('ID="1" Pd="20"', 'ID="1" Balance="2" Pd="20"'),
                      XML.replace('Time="11:10:00"', 'Time="13:10:00"'),
                      XML.replace('<doc Ver="1">', '<!DOCTYPE doc><doc Ver="1">')):
            with self.subTest(value=value[:80]):
                self.path.write_text(value, encoding="utf-8")
                with self.assertRaises(JwgbNativeParseError):
                    load_file(self.path)

    def test_reconstructed_bet_reaches_default_import_and_exports_with_provenance(self):
        from openpyxl import load_workbook
        from tristar_bet.ui.main_window import export_results_xlsx, active_metric_rows
        self.path.write_text(reconstructable_xml(), encoding='utf-8')
        result = load_file(self.path)
        fit = bet_analysis(result)
        self.assertTrue(fit.ok)
        self.assertGreater(fit.surface_area_m2_g, 0)
        self.assertEqual(fit.status, 'warning_reconstructed_monolayer_outside_range')
        self.assertEqual((fit.pressure_min, fit.pressure_max), (.08, .3))
        self.assertEqual(result.point_count, 4)
        self.assertTrue(result.method_options['jwgb_reconstruction_closure_applied'])
        self.assertFalse(result.method_options['jwgb_reconstruction_nonideality_applied'])
        self.assertEqual(result.method_options['jwgb_reconstruction_validation'],
                         'bk200c_manual_q_one_same_run_export_match')
        self.assertIn('尚未与本文件原厂结果核验', str(dict(active_metric_rows(result))['多点 BET 状态']))
        self.assertIn('单层容量超出选区', str(dict(active_metric_rows(result))['多点 BET 状态']))
        outputs = export_results_csv([result], Path(self.temp.name) / 'csv')
        csv_path = next(p for p in outputs if p.name.endswith('_isotherm.csv'))
        with csv_path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        self.assertTrue(all(r['quantity_status'] == 'reconstructed_unvalidated' for r in rows))
        summary_path = next(p for p in outputs if p.name.endswith('_summary.json'))
        self.assertEqual(json.loads(summary_path.read_text(encoding='utf-8'))[0]['quantity_status'], 'reconstructed_unvalidated')
        path = Path(self.temp.name) / 'export.xlsx'
        export_results_xlsx([result], path)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in ('摘要', '实际等温线', 'BET'):
                values = [str(v) for row in wb[sheet].iter_rows(values_only=True) for v in row]
                self.assertTrue(any('尚未与本文件原厂结果核验' in v for v in values), sheet)
        finally:
            wb.close()

    def test_reconstruction_fails_atomically_for_uncovered_or_corrupt_calibration(self):
        changes = [('<AutoTestQ>0', '<AutoTestQ>1'), ('00.01.05.00', '99.00.00.00'),
                   ('<QValue>0.0085', '<QValue>NaN'), ('PcdT="26.85"', 'PcdT="NaN"'),
                   ('<IsRealGas>0', '<IsRealGas>1')]
        for old, new in changes:
            with self.subTest(field=old):
                self.path.write_text(reconstructable_xml().replace(old, new), encoding='utf-8')
                result = load_file(self.path)
                self.assertTrue(all(p.quantity_adsorbed_cm3_g_stp is None for p in result.isotherm))
                self.assertNotIn('jwgb_reconstructed', result.method_options)
                self.assertIsNone(bet_analysis(result).surface_area_m2_g)

    def test_auto_q_valid_synthetic_run_computes_bet_using_unrounded_calibration(self):
        # Independent ideal helium expansion: 24 -> 15 kPa at 300 K.
        q = (273.15 / 101.325) * (24 - 15) / (300 * 15)
        calibration = '''<Q>
<Item Pd="-0.00001" Pcd="0" PdT="0" PcdT="26.85" Day="2026-10-12" Time="11:00:10"/>
<Item Pd="24" Pcd="15" PdT="26.85" PcdT="26.85" Day="2026-10-12" Time="11:00:20"/>
</Q>'''
        xml = (reconstructable_xml().replace('<AutoTestQ>0', '<AutoTestQ>1')
               .replace('<QValue>0.0085', f'<QValue>{q:.6f}')
               .replace('</BaseInfo>', '</BaseInfo>' + calibration))
        self.path.write_text(xml, encoding='utf-8')
        result = load_file(self.path)
        self.assertAlmostEqual(result.method_options['jwgb_reconstruction_q'], q, places=14)
        self.assertNotEqual(result.method_options['jwgb_reconstruction_q'], float(f'{q:.6f}'))
        self.assertEqual(result.method_options['jwgb_reconstruction_calibration_source'], 'helium_calibration_records')
        self.assertGreater(bet_analysis(result).surface_area_m2_g, 0)
        self.assertEqual(result.method_options['quantity_status'], 'reconstructed_unvalidated')

    def test_reconstructed_monolayer_warning_clears_for_a_bracketing_range(self):
        from dataclasses import replace
        self.path.write_text(reconstructable_xml(), encoding='utf-8')
        result = load_file(self.path)
        # Known BET curve with Vm=27 and C=100, whose monolayer lies at p/p0~.091.
        points = []
        for point, x in zip(result.isotherm[:3], (.08, .15, .25)):
            v = 27 * 100 * x / ((1 - x) * (1 + 99 * x))
            points.append(replace(point, relative_pressure=x, quantity_adsorbed_cm3_g_stp=v))
        result = replace(result, isotherm=points)
        self.assertEqual(bet_analysis(result).status, 'warning_reconstructed')
        # Changing only the selection excludes the known monolayer.
        shifted = [replace(p, relative_pressure=x,
                           quantity_adsorbed_cm3_g_stp=27 * 100 * x / ((1 - x) * (1 + 99 * x)))
                   for p, x in zip(points, (.15, .20, .25))]
        result = replace(result, isotherm=shifted)
        self.assertEqual(bet_analysis(result).status, 'warning_reconstructed_monolayer_outside_range')

    def test_negative_reconstruction_is_retained_and_not_silently_filtered_for_bet(self):
        self.path.write_text(reconstructable_xml().replace('<QValue>0.0085', '<QValue>0.02'), encoding='utf-8')
        result = load_file(self.path)
        self.assertLess(result.isotherm[0].quantity_adsorbed_cm3_g_stp, 0)
        fit = bet_analysis(result)
        self.assertEqual(fit.status, 'nonpositive_reconstructed_uptake')
        self.assertIsNone(fit.surface_area_m2_g)

    def test_ui_import_summary_and_plots(self):
        from PyQt5 import QtCore, QtWidgets
        from tristar_bet.ui.main_window import FileImportDialog, MainWindow, summary_rows
        app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
        dialog = FileImportDialog(initial_dir=self.path.parent)
        self.assertEqual(dialog.available_table.rowCount(), 1)
        dialog.close()
        dialog.deleteLater()
        settings = QtCore.QSettings(str(Path(self.temp.name) / 'settings.ini'), QtCore.QSettings.IniFormat)
        with patch.object(MainWindow, '_auto_check_for_updates'), patch.object(MainWindow, '_initialize_telemetry'), patch('tristar_bet.ui.main_window.QtCore.QSettings', return_value=settings):
            window = MainWindow()
        try:
            window.load_files([str(self.path)], replace=True)
            self.assertEqual(len(window.results), 1)
            self.assertNotIn('导入状态', dict(summary_rows(window.results[0])))
            self.assertEqual(window.results[0].method_options['quantity_status'], 'unavailable_native_pressure_only')
            from tristar_bet.ui.main_window import BET_COLUMN
            self.assertIn('尚未解析出吸附量', window.sample_table.item(0, BET_COLUMN).toolTip())
            for index in range(window.plot_tabs.count()):
                window.plot_tabs.setCurrentIndex(index)
                app.processEvents()
            window.refresh_all()
            # A supported profile must reach the same real UI path, not only a probe.
            self.path.write_text(reconstructable_xml(), encoding='utf-8')
            window.load_files([str(self.path)], replace=True)
            from tristar_bet.ui.main_window import BET_COLUMN
            self.assertNotIn('重建', window.sample_table.item(0, BET_COLUMN).text())
            self.assertEqual(window.sample_table.item(0, BET_COLUMN).toolTip(), '区间计算完成')
            self.assertNotIn('导入状态', dict(summary_rows(window.results[0])))
            self.assertEqual(window.isotherm_plot.getPlotItem().titleLabel.text, '吸附/脱附等温线')
            window._refresh_sample_bet_cell(0)
            self.assertNotIn('重建', window.sample_table.item(0, BET_COLUMN).text())
            self.assertEqual(window.sample_table.item(0, BET_COLUMN).toolTip(), '区间计算完成')
            for index in range(window.plot_tabs.count()):
                window.plot_tabs.setCurrentIndex(index)
                app.processEvents()
        finally:
            window.close()
            window.deleteLater()
            QtCore.QCoreApplication.sendPostedEvents(None, QtCore.QEvent.DeferredDelete)


if __name__ == '__main__':
    unittest.main()
