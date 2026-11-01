from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PyQt5 import QtCore, QtWidgets

from tristar_bet.models import IsothermPoint
from tristar_bet.ui.main_window import MainWindow, export_results_xlsx
from test_detail_result_tables import sample

KJS = {'numerator': 60.65, 'offset': 0.03071, 'exponent': 0.3968}


def with_desorption(result):
    ads = list(result.isotherm)
    des = [IsothermPoint(len(ads) + i + 1, 'desorption', 0, p.absolute_pressure_mmHg, p.relative_pressure,
                         p.raw_internal_cm3_stp * 1.02, 760, 3600 + i * 60,
                         p.quantity_adsorbed_cm3_g_stp * 1.02, p.quantity_adsorbed_mmol_g * 1.02)
           for i, p in enumerate(reversed(ads[:-1]))]
    return result.__class__(**{**result.__dict__, 'isotherm': ads + des})


class PoreDistributionExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        settings = QtCore.QSettings(str(Path(self.temp.name) / 's.ini'), QtCore.QSettings.IniFormat)
        with patch.object(MainWindow, '_auto_check_for_updates'), patch.object(MainWindow, '_initialize_telemetry'), \
                patch('tristar_bet.ui.main_window.QtCore.QSettings', return_value=settings):
            self.w = MainWindow()
        self.samples = [with_desorption(sample('A', 1)), with_desorption(sample('B', 2))]
        self.w.results = list(self.samples)
        self.w.visible_results = [True, True]
        self.w.active_index = 0
        self.w.refresh_isotherm_plot = Mock()
        self.w.refresh_analysis_plots = Mock()
        for result in self.samples:  # portable thickness model; both BJH branches on
            self.w.custom_bjh_settings[id(result)] = {'thickness_method': 'kjs', 'show_adsorption': True,
                                                      'show_desorption': True}
            self.w.custom_dh_settings[id(result)] = {'thickness_method': 'kjs'}

    def tearDown(self):
        self.w.close()
        self.w.deleteLater()
        QtCore.QCoreApplication.sendPostedEvents(None, QtCore.QEvent.DeferredDelete)
        self.temp.cleanup()

    def export(self):
        from openpyxl import load_workbook
        path = Path(self.temp.name) / 'out.xlsx'
        export_results_xlsx(self.samples, path,
                            pore_distributions=[self.w._pore_distribution_exports(r) for r in self.samples])
        return load_workbook(BytesIO(path.read_bytes()), data_only=True)

    def blocks(self, sheet):
        return [(c, sheet.cell(1, c).value) for c in range(1, sheet.max_column + 1) if sheet.cell(1, c).value]

    def test_four_method_sheets_with_side_by_side_blocks(self):
        wb = self.export()
        for name in ('BET', 'Langmuir', 't-Plot', 'BJH', 'DH', 'HK', 'DFT'):
            self.assertIn(name, wb.sheetnames)
        bjh = wb['BJH']
        labels = [label for _c, label in self.blocks(bjh)]
        self.assertEqual(labels, ['A.SMP | A | 吸附支', 'A.SMP | A | 脱附支', 'B.SMP | B | 吸附支', 'B.SMP | B | 脱附支'])
        self.assertEqual(len(self.blocks(wb['HK'])), 2)
        self.assertEqual(len(self.blocks(wb['DFT'])), 4)  # distribution + fit per sample

    def test_bjh_values_equal_what_the_window_computes(self):
        wb = self.export()
        sheet = wb['BJH']
        column = self.blocks(sheet)[0][0]
        headers = [sheet.cell(3, column + i).value for i in range(6)]
        self.assertEqual(headers[:2], ['平均孔径 D (nm)', '孔径下限 (nm)'])
        self.assertIn('dV/dlogD (cm3/g)', headers)
        rows = self.w._cached_bjh_distribution_rows(self.samples[0], phase='adsorption', thickness_method='kjs',
                                                    thickness_params=KJS, correction='standard',
                                                    open_pore_fraction=0.0, smooth=True)
        self.assertGreater(len(rows), 3)
        exported = [sheet.cell(4 + i, column).value for i in range(len(rows))]
        for value, row in zip(exported, rows):  # xlsx round-trip may change the last binary digit
            self.assertAlmostEqual(value, row['pore_diameter_nm'], places=12)
        dlog = headers.index('dV/dlogD (cm3/g)')
        self.assertAlmostEqual(sheet.cell(4, column + dlog).value, rows[0]['differential_pore_volume_cm3_g'])
        self.assertIn('kjs', sheet.cell(2, column).value)

    def test_hk_block_is_plottable_numbers(self):
        sheet = self.export()['HK']
        column = self.blocks(sheet)[0][0]
        self.assertEqual(sheet.cell(3, column).value, '孔宽 W (nm)')
        values = [sheet.cell(r, column).value for r in range(4, sheet.max_row + 1) if sheet.cell(r, column).value is not None]
        self.assertTrue(values and all(isinstance(v, float) for v in values))

    def test_failing_method_is_reported_without_breaking_export(self):
        with patch.object(self.w, '_cached_dft_result', side_effect=RuntimeError('kernel missing')):
            wb = self.export()
        sheet = wb['DFT']
        self.assertIn('计算失败', sheet.cell(2, 1).value)
        self.assertIn('kernel missing', sheet.cell(2, 1).value)
        self.assertIn('BJH', wb.sheetnames)

    def test_branch_switched_off_is_not_exported(self):
        for result in self.samples:
            self.w.custom_bjh_settings[id(result)]['show_desorption'] = False
        labels = [label for _c, label in self.blocks(self.export()['BJH'])]
        self.assertTrue(all('吸附支' in label for label in labels))

    def test_without_distributions_export_is_unchanged(self):
        from openpyxl import load_workbook
        path = Path(self.temp.name) / 'plain.xlsx'
        export_results_xlsx(self.samples, path)
        self.assertNotIn('BJH', load_workbook(BytesIO(path.read_bytes())).sheetnames)


if __name__ == '__main__':
    unittest.main()
